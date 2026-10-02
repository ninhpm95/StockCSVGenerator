"""Everything between a file on disk and a resolved column map: CSV / XLSX
reading (step 010), holdings-marker trimming (020), header-row detection (030)
and column resolution (040). Inputs are opened read-only; nothing is ever
written back to them."""
from __future__ import annotations

import re
import csv
from datetime import date, datetime
from io import StringIO
from pathlib import Path

from .config import (CODE_COLUMN_CANDIDATES, COLUMN_RESOLUTION_ORDER,
                    EXCHANGE_COLUMN_CANDIDATES, HEADER_KEYWORD_COMBINATIONS,
                    HOLDINGS_MARK, HOLDINGS_SHEET, ISIN_COLUMN_CANDIDATES,
                    NAME_COLUMN_CANDIDATES, PRICE_COLUMN_CANDIDATES,
                    SHARES_COLUMN_CANDIDATES,
                    VALUATION_COLUMN_CANDIDATES, WEIGHT_COLUMN_CANDIDATES)
from .models import ColumnMap, Grid

SUPPORTED_SUFFIXES = {".csv", ".xlsx", ".xlsm"}


# ---------------------------------------------------------------- reading --
def load_file(path: Path) -> Grid:
    """One file -> one Grid (a CSV, or the holdings sheet of a workbook)."""
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"unsupported file type: {suffix}")
    return read_csv(path) if suffix == ".csv" else read_xlsx(path)


def read_csv(path: Path) -> Grid:
    text = decode(path.read_bytes())
    rows = [[c.strip() for c in row]
            for row in csv.reader(StringIO(text), delimiter=sniff_delimiter(text))]
    return Grid(rows=pad(drop_empty_rows(rows)))


def read_xlsx(path: Path) -> Grid:
    """Only one holdings sheet is read; the rest of the workbook is ignored.
    HOLDINGS_SHEET is tried left to right: the first name present in the
    workbook wins and the remaining names are not looked at."""
    from openpyxl import load_workbook    # lazy: csv-only runs don't need it
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_name = next((s for s in HOLDINGS_SHEET if s in workbook.sheetnames), None)
        if sheet_name is None:
            wanted = ", ".join(f"'{s}'" for s in HOLDINGS_SHEET)
            raise ValueError(f"no holdings sheet found (looked for {wanted}; "
                             f"found: {', '.join(workbook.sheetnames)})")
        ws = workbook[sheet_name]
        rows = [[cell_to_str(getattr(c, "value", None), getattr(c, "number_format", None)) for c in row] for row in ws.iter_rows()]
    finally:
        workbook.close()
    rows = drop_empty_rows([[c.strip() for c in r] for r in rows])
    return Grid(rows=pad(rows))


def decode(raw: bytes) -> str:
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw.decode("utf-8-sig")
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    for encoding in ("utf-8", "cp932"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


SNIFF_LINES = 10


def sniff_delimiter(text: str) -> str:
    """Delimiter with the highest total count over the first SNIFF_LINES
    non-empty lines (a single title line can't mislead it); ',' if none."""
    counts = {d: 0 for d in (",", ";", "\t")}
    seen = 0
    for line in text.splitlines():
        if not line.strip():
            continue
        for d in counts:
            counts[d] += line.count(d)
        seen += 1
        if seen == SNIFF_LINES:
            break
    best = max(counts, key=counts.get)
    return best if counts[best] else ","


def drop_empty_rows(rows: list[list[str]]) -> list[list[str]]:
    return [r for r in rows if any(c for c in r)]


def pad(rows: list[list[str]]) -> list[list[str]]:
    width = max((len(r) for r in rows), default=0)
    return [r + [""] * (width - len(r)) for r in rows]


_QUOTED = re.compile(r'"[^"]*"|\\.')


def _is_percent_format(number_format) -> bool:
    """True for Excel formats like 0%, 0.00%. A '%' inside quotes or escaped
    (0.0"%") is just a label and does not scale the value."""
    return bool(number_format) and "%" in _QUOTED.sub("", str(number_format))


def cell_to_str(value, number_format=None) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)) and not isinstance(value, bool) \
            and _is_percent_format(number_format):
        # Excel stores 80% as 0.8; give back the number the cell displays.
        pct = f"{round(float(value) * 100, 10):.10f}".rstrip("0").rstrip(".")
        return pct + "%"
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return repr(value)
    return str(value)


def row_matches(row: list[str], keywords: tuple[str, ...]) -> bool:
    """All keywords found, in order, scanning cells left to right (case sensitive)."""
    ki = 0
    for cell in row:
        if ki == len(keywords):
            return True
        if keywords[ki] in cell:
            ki += 1
    return ki == len(keywords)


def find_last_marker_row(grid: Grid) -> int | None:
    """Index of the LAST row with HOLDINGS_MARK in any cell, or None."""
    for i in range(len(grid.rows) - 1, -1, -1):
        if any(HOLDINGS_MARK in cell for cell in grid.rows[i]):
            return i
    return None


def find_header_row(grid: Grid) -> int | None:
    """Index of the first row matching any keyword combination (tuple order)."""
    for i, row in enumerate(grid.rows):
        for keywords in HEADER_KEYWORD_COMBINATIONS:
            if row_matches(row, keywords):
                return i
    return None


_CANDIDATES = {
    "code": CODE_COLUMN_CANDIDATES,
    "name": NAME_COLUMN_CANDIDATES,
    "isin": ISIN_COLUMN_CANDIDATES,
    "shares": SHARES_COLUMN_CANDIDATES,
    "price": PRICE_COLUMN_CANDIDATES,
    "valuation": VALUATION_COLUMN_CANDIDATES,
    "weight": WEIGHT_COLUMN_CANDIDATES,
    "exchange": EXCHANGE_COLUMN_CANDIDATES,
}
# Order comes from config.COLUMN_RESOLUTION_ORDER.
_COLUMN_GROUPS = tuple((f, _CANDIDATES[f]) for f in COLUMN_RESOLUTION_ORDER)


def resolve_columns(header: list[str]) -> ColumnMap:
    """Case-sensitive match of candidates against the raw header cells (no
    normalisation). Fields claim columns in COLUMN_RESOLUTION_ORDER. Within a
    field:
      1. an EXACT match (cell == candidate) wins; candidates are tried in list
         order and the leftmost untaken cell is used;
      2. only if no candidate matches exactly, the first candidate (list order)
         that is a SUBSTRING of an untaken cell wins, leftmost cell first."""
    taken: set[int] = set()
    found: dict[str, int] = {}
    for field_name, candidates in _COLUMN_GROUPS:
        idx = None
        for candidate in candidates:                      # pass 1: exact
            idx = next((i for i, cell in enumerate(header)
                        if i not in taken and cell == candidate), None)
            if idx is not None:
                break
        if idx is None:
            for candidate in candidates:                  # pass 2: substring
                idx = next((i for i, cell in enumerate(header)
                            if i not in taken and candidate in cell), None)
                if idx is not None:
                    break
        if idx is not None:
            found[field_name] = idx
            taken.add(idx)
    return ColumnMap(**found)
